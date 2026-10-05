#!/usr/bin/env python3
"""Alt taşeron sözleşmesi Markdown → Word (.docx)."""

from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Cm, Pt, RGBColor

ROOT = Path(__file__).resolve().parents[1]
MD = ROOT / "taseron_paket" / "Alt_Taseron_Demir_Kalip_Sozlesmesi_Gundogdu_KM.md"
OUT = ROOT / "taseron_paket" / "Alt_Taseron_Demir_Kalip_Sozlesmesi_Gundogdu_KM.docx"


def set_run_font(run, size=11, bold=False, italic=False, color=None):
    run.font.name = "Calibri"
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = italic
    if color:
        run.font.color.rgb = color


def add_table(doc, rows: list[list[str]]):
    if not rows:
        return
    ncols = max(len(r) for r in rows)
    table = doc.add_table(rows=len(rows), cols=ncols)
    table.style = "Table Grid"
    for i, row in enumerate(rows):
        for j in range(ncols):
            table.rows[i].cells[j].text = row[j] if j < len(row) else ""


def parse_md_to_docx():
    doc = Document()
    for section in doc.sections:
        section.top_margin = Cm(2)
        section.bottom_margin = Cm(2)
        section.left_margin = Cm(2.5)
        section.right_margin = Cm(2.5)

    lines = MD.read_text(encoding="utf-8").splitlines()
    table_buf: list[str] = []

    def flush_table():
        nonlocal table_buf
        if not table_buf:
            return
        parsed: list[list[str]] = []
        for r in table_buf:
            if all(set(c.strip()) <= set("-:") for c in r.split("|") if c.strip()):
                continue
            parsed.append([c.strip().replace("**", "") for c in r.strip().strip("|").split("|")])
        add_table(doc, parsed)
        doc.add_paragraph()
        table_buf = []

    for line in lines:
        if line.startswith("|") and line.count("|") >= 2:
            table_buf.append(line)
            continue
        flush_table()

        if line.startswith("# "):
            t = doc.add_heading(line[2:].strip(), level=0)
            t.alignment = WD_ALIGN_PARAGRAPH.CENTER
        elif line.startswith("## "):
            doc.add_page_break()
            doc.add_heading(line[3:].strip(), level=1)
        elif line.startswith("### "):
            doc.add_heading(line[4:].strip(), level=2)
        elif line.startswith("> "):
            p = doc.add_paragraph()
            set_run_font(p.add_run(line[2:].strip()), size=10, italic=True, color=RGBColor(0x44, 0x44, 0x44))
        elif line.strip() == "---":
            doc.add_paragraph()
        elif line.strip().startswith("- [ ]"):
            doc.add_paragraph("☐ " + line.strip()[5:].strip(), style="List Bullet")
        elif line.strip().startswith("- "):
            doc.add_paragraph(line.strip()[2:].replace("**", ""), style="List Bullet")
        elif not line.strip():
            continue
        else:
            p = doc.add_paragraph(line.replace("**", ""))
            for run in p.runs:
                set_run_font(run)

    flush_table()
    doc.save(OUT)
    print(f"OK: {OUT}")


def docx_to_pdf():
    import subprocess

    pdf = OUT.with_suffix(".pdf")
    subprocess.run(
        [
            "libreoffice",
            "--headless",
            "--convert-to",
            "pdf",
            "--outdir",
            str(OUT.parent),
            str(OUT),
        ],
        check=True,
        capture_output=True,
    )
    print(f"OK: {pdf}")


if __name__ == "__main__":
    parse_md_to_docx()
    try:
        docx_to_pdf()
    except (FileNotFoundError, subprocess.CalledProcessError) as e:
        print("PDF atlandı (LibreOffice gerekli):", e)
