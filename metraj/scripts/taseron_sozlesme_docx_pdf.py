#!/usr/bin/env python3
"""Taşeronlar arası kısa sözleşme → DOCX + PDF."""

import subprocess
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Cm, Pt, RGBColor

ROOT = Path(__file__).resolve().parents[1]
MD = ROOT / "taseron_paket" / "Taseron_Demir_Kalip_Sozlesmesi_Basit.md"
OUT = ROOT / "taseron_paket" / "Taseron_Demir_Kalip_Sozlesmesi_Basit.docx"


def md_to_docx():
    doc = Document()
    for section in doc.sections:
        section.top_margin = Cm(2)
        section.bottom_margin = Cm(2)
        section.left_margin = Cm(2.5)
        section.right_margin = Cm(2.5)

    table_buf: list[str] = []

    def flush_table():
        nonlocal table_buf
        if not table_buf:
            return
        rows: list[list[str]] = []
        for r in table_buf:
            if all(set(c.strip()) <= set("-:") for c in r.split("|") if c.strip()):
                continue
            rows.append([c.strip().replace("**", "") for c in r.strip().strip("|").split("|")])
        if rows:
            ncols = max(len(x) for x in rows)
            t = doc.add_table(rows=len(rows), cols=ncols)
            t.style = "Table Grid"
            for i, row in enumerate(rows):
                for j in range(ncols):
                    t.rows[i].cells[j].text = row[j] if j < len(row) else ""
            doc.add_paragraph()
        table_buf = []

    for line in MD.read_text(encoding="utf-8").splitlines():
        if line.startswith("|") and line.count("|") >= 2:
            table_buf.append(line)
            continue
        flush_table()
        if line.startswith("# "):
            h = doc.add_heading(line[2:].strip(), level=0)
            h.alignment = WD_ALIGN_PARAGRAPH.CENTER
        elif line.startswith("## "):
            doc.add_heading(line[3:].strip(), level=1)
        elif line.strip() == "---":
            doc.add_paragraph()
        elif line.strip().startswith("- "):
            doc.add_paragraph(line.strip()[2:], style="List Bullet")
        elif line.strip():
            p = doc.add_paragraph(line.replace("**", ""))
            for run in p.runs:
                run.font.name = "Calibri"
                run.font.size = Pt(11)
    flush_table()
    doc.save(OUT)


def to_pdf():
    subprocess.run(
        ["libreoffice", "--headless", "--convert-to", "pdf", "--outdir", str(OUT.parent), str(OUT)],
        check=True,
        capture_output=True,
    )


if __name__ == "__main__":
    md_to_docx()
    to_pdf()
    print(OUT)
    print(OUT.with_suffix(".pdf"))
