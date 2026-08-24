#!/usr/bin/env python3
"""Götürü kalıp hakediş tablosu — Excel üretici."""

from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "Hakedis_Tablosu_Goturu_9400.xlsx"

M2 = 9400
BF = 500
DONEMLER = [
    ("1", "Bodrum + temel", 1800),
    ("2", "Zemin + 1. kat", 2400),
    ("3", "2. + 3. kat", 2800),
    ("4", "Çatı + söküm", 2400),
]


def money(n: float) -> float:
    return round(n, 2)


def build():
    wb = Workbook()
    ws = wb.active
    ws.title = "Hakediş"

    header_fill = PatternFill("solid", fgColor="1E293B")
    header_font = Font(bold=True, color="FFFFFF", size=10)
    title_font = Font(bold=True, size=14, color="B45309")

    ws["A1"] = "GÖTÜRÜ KALIP HAKEDİŞ TABLOSU"
    ws["A1"].font = title_font
    ws.merge_cells("A1:J1")

    meta = [
        ("Proje", "Karşıyaka Ortaokulu — Kalıp işçiliği"),
        ("Metraj", f"Düz ölçü — {M2:,.0f} m²".replace(",", ".")),
        ("Birim fiyat", f"{BF} TL/m² (KDV hariç)"),
        ("Götürü matrah", f"{M2 * BF:,.0f} TL".replace(",", ".")),
        ("KDV / Tevkifat", "%20 KDV · 4/10 tevkifat"),
        ("SGK", "Ana firma tarafından"),
    ]
    row = 3
    for label, val in meta:
        ws.cell(row=row, column=1, value=label).font = Font(bold=True)
        ws.cell(row=row, column=2, value=val)
        row += 1

    row += 1
    headers = [
        "Dönem",
        "Açıklama",
        "Tamamlanan (m²)",
        "Kümülatif (m²)",
        "BF (TL/m²)",
        "Matrah (KDV hariç)",
        "KDV %20",
        "Tevkifat 4/10",
        "Tahsil (matrah+6/10 KDV)",
        "Küm. matrah",
        "Küm. tahsil",
    ]
    for col, h in enumerate(headers, 1):
        c = ws.cell(row=row, column=col, value=h)
        c.font = header_font
        c.fill = header_fill
        c.alignment = Alignment(horizontal="center", wrap_text=True)

    kum_m2 = kum_matrah = kum_tahsil = 0.0
    row += 1
    for no, aciklama, m2 in DONEMLER:
        kum_m2 += m2
        matrah = m2 * BF
        kdv = money(matrah * 0.20)
        tevkifat = money(kdv * 0.40)
        tahsil = money(matrah + kdv - tevkifat)
        kum_matrah += matrah
        kum_tahsil += tahsil

        vals = [no, aciklama, m2, kum_m2, BF, matrah, kdv, tevkifat, tahsil, kum_matrah, kum_tahsil]
        for col, v in enumerate(vals, 1):
            c = ws.cell(row=row, column=col, value=v)
            if col >= 3:
                c.number_format = "#,##0.00"
        row += 1

    row += 1
    ws.cell(row=row, column=2, value="TOPLAM").font = Font(bold=True)
    for col, v in [(3, M2), (6, M2 * BF), (7, M2 * BF * 0.2), (8, M2 * BF * 0.2 * 0.4), (9, kum_tahsil)]:
        c = ws.cell(row=row, column=col, value=money(v) if col > 3 else v)
        c.font = Font(bold=True)
        c.number_format = "#,##0.00"

    for col, width in enumerate([8, 22, 14, 14, 12, 16, 12, 14, 18, 14, 14], 1):
        ws.column_dimensions[chr(64 + col)].width = width

    wb.save(OUT)
    print(f"Yazıldı: {OUT}")


if __name__ == "__main__":
    build()
