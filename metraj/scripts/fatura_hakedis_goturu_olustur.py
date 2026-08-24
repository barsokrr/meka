#!/usr/bin/env python3
"""Götürü kalıp — 4 dönem hakediş + e-Fatura satır şablonu üretici."""

from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

ROOT = Path(__file__).resolve().parents[1]

M2 = 9400
BF = 500
PROJE = "Karşıyaka Ortaokulu (24 derslik) — Kalıp işçiliği"
PROJE_NO = "MEBİZ.73-10-25-01-SU-001-R0"
SATICI = "ABDURRAHMAN BARIŞ ÖKER"
VKN = "6530560679"
VD = "Van Vergi Dairesi Müdürlüğü"
ADRES = "Yeni Mah. Çalıbaşı Van İpekyolu No:26/1, Van"
NACE = "43.99.05"

DONEMLER = [
    ("1", "Bodrum + temel kalıp işçiliği", 1800),
    ("2", "Zemin + 1. kat kalıp işçiliği", 2400),
    ("3", "2. + 3. kat kalıp işçiliği", 2800),
    ("4", "Çatı katı + söküm kalıp işçiliği", 2400),
]


def money(n: float) -> float:
    return round(n, 2)


def calc(m2: float) -> dict:
    matrah = money(m2 * BF)
    kdv = money(matrah * 0.20)
    tevkifat = money(kdv * 0.40)
    tahsil = money(matrah + kdv - tevkifat)
    return {"matrah": matrah, "kdv": kdv, "tevkifat": tevkifat, "tahsil": tahsil}


def thin_border():
    s = Side(style="thin", color="CBD5E1")
    return Border(left=s, right=s, top=s, bottom=s)


def write_header(ws, row: int, headers: list):
    fill = PatternFill("solid", fgColor="1E293B")
    font = Font(bold=True, color="FFFFFF", size=10)
    for col, h in enumerate(headers, 1):
        c = ws.cell(row=row, column=col, value=h)
        c.font = font
        c.fill = fill
        c.alignment = Alignment(horizontal="center", wrap_text=True)
        c.border = thin_border()


def build_kapak(wb: Workbook):
    ws = wb.active
    ws.title = "Kapak"
    ws["A1"] = "HAKEDİŞ & FATURA PAKETİ"
    ws["A1"].font = Font(bold=True, size=16, color="B45309")
    meta = [
        ("Proje", PROJE),
        ("Proje no", PROJE_NO),
        ("Metraj yöntemi", f"Düz ölçü — {M2:,} m²".replace(",", ".")),
        ("Götürü birim fiyat", f"{BF} TL/m² (KDV hariç)"),
        ("Taşeron / satıcı", SATICI),
        ("VKN", VKN),
        ("Vergi dairesi", VD),
        ("NACE", NACE),
        ("KDV tevkifat", "Yapım işi — 4/10"),
        ("SGK", "Ana yüklenici tarafından"),
        ("Alıcı (doldurun)", "___________________________"),
        ("Alıcı VKN (doldurun)", "___________________________"),
    ]
    r = 3
    for k, v in meta:
        ws.cell(row=r, column=1, value=k).font = Font(bold=True)
        ws.cell(row=r, column=2, value=v)
        r += 1
    ws.column_dimensions["A"].width = 22
    ws.column_dimensions["B"].width = 55


def build_ozet(wb: Workbook):
    ws = wb.create_sheet("Hakediş Özet")
    headers = [
        "Dönem",
        "Açıklama",
        "m²",
        "Matrah",
        "KDV %20",
        "Tevkifat 4/10",
        "Tahsil",
        "Fatura no",
        "Fatura tarihi",
        "Ödendi",
    ]
    write_header(ws, 1, headers)
    kum_matrah = kum_tahsil = 0.0
    row = 2
    for no, aciklama, m2 in DONEMLER:
        h = calc(m2)
        kum_matrah += h["matrah"]
        kum_tahsil += h["tahsil"]
        vals = [no, aciklama, m2, h["matrah"], h["kdv"], h["tevkifat"], h["tahsil"], "", "", ""]
        for col, v in enumerate(vals, 1):
            c = ws.cell(row=row, column=col, value=v)
            c.border = thin_border()
            if col >= 4:
                c.number_format = "#,##0.00"
        row += 1
    ws.cell(row=row, column=2, value="TOPLAM").font = Font(bold=True)
    for col, v in [(3, M2), (4, kum_matrah), (5, kum_matrah * 0.2), (6, kum_matrah * 0.2 * 0.4), (7, kum_tahsil)]:
        c = ws.cell(row=row, column=col, value=money(v) if col > 3 else v)
        c.font = Font(bold=True)
        c.number_format = "#,##0.00"
    for col, w in enumerate([8, 28, 10, 14, 12, 14, 14, 14, 14, 10], 1):
        ws.column_dimensions[chr(64 + col)].width = w


def build_donem_sheet(wb: Workbook, no: str, aciklama: str, m2: float):
    ws = wb.create_sheet(f"Hakediş {no}")
    h = calc(m2)
    ws["A1"] = f"{no}. HAKEDİŞ — METRAJ TUTANAĞI & FATURA ÖZETİ"
    ws["A1"].font = Font(bold=True, size=13, color="B45309")

    lines = [
        ("Proje", PROJE),
        ("Dönem", f"{no}. {aciklama}"),
        ("Tamamlanan metraj", f"{m2:,.1f} m²".replace(",", ".")),
        ("Birim fiyat", f"{BF} TL/m² (KDV hariç)"),
        ("Matrah", f"{h['matrah']:,.2f} TL".replace(",", ".")),
        ("KDV %20", f"{h['kdv']:,.2f} TL".replace(",", ".")),
        ("Tevkifat 4/10", f"{h['tevkifat']:,.2f} TL".replace(",", ".")),
        ("Tahsil edilecek", f"{h['tahsil']:,.2f} TL".replace(",", ".")),
    ]
    r = 3
    for k, v in lines:
        ws.cell(row=r, column=1, value=k).font = Font(bold=True)
        ws.cell(row=r, column=2, value=v)
        r += 1

    r += 2
    ws.cell(row=r, column=1, value="Fatura açıklama metni:").font = Font(bold=True)
    r += 1
    fatura_text = (
        f"{PROJE} — {no}. hakediş — {aciklama} — "
        f"{m2:,.1f} m² × {BF} TL/m² = {h['matrah']:,.2f} TL (KDV hariç). "
        f"Götürü kalıp işçiliği; malzeme hariç; SGK ana yüklenici. NACE {NACE}. KDV tevkifat 4/10."
    ).replace(",", ".")
    ws.cell(row=r, column=1, value=fatura_text)
    ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=4)
    ws.cell(row=r, column=1).alignment = Alignment(wrap_text=True)

    r += 3
    write_header(
        ws,
        r,
        ["Sıra", "Kod", "Mal/Hizmet açıklaması", "Birim", "Miktar", "BF", "Matrah", "KDV", "Tevkifat", "Tahsil"],
    )
    r += 1
    desc = f"Götürü kalıp işçiliği — {no}. hakediş — {aciklama} (malzeme hariç)"
    vals = [1, "43.99.05 / 15.180.1002", desc, "m²", m2, BF, h["matrah"], h["kdv"], h["tevkifat"], h["tahsil"]]
    for col, v in enumerate(vals, 1):
        c = ws.cell(row=r, column=col, value=v)
        c.border = thin_border()
        if col >= 5:
            c.number_format = "#,##0.00"

    r += 4
    ws.cell(row=r, column=1, value="Taşeron imza:").font = Font(bold=True)
    ws.cell(row=r, column=3, value="İşveren / şantiye onayı:").font = Font(bold=True)
    r += 1
    ws.cell(row=r, column=1, value="Tarih: ___ / ___ / 20__")
    ws.cell(row=r, column=3, value="Tarih: ___ / ___ / 20__")

    for col, w in enumerate([6, 18, 40, 8, 10, 10, 14, 12, 12, 14], 1):
        ws.column_dimensions[chr(64 + col)].width = w


def export_csv():
    out_dir = ROOT
    for no, aciklama, m2 in DONEMLER:
        h = calc(m2)
        path = out_dir / f"Fatura_Hakedis_Donem_{no}.csv"
        lines = [
            "Alan,Deger",
            f"Donem,{no}",
            f"Aciklama,{aciklama}",
            f"Satici,{SATICI}",
            f"VKN,{VKN}",
            f"Proje,{PROJE}",
            f"Miktar_m2,{m2}",
            f"Birim_Fiyat,{BF}",
            f"Matrah,{h['matrah']}",
            f"KDV_20,{h['kdv']}",
            f"Tevkifat_4_10,{h['tevkifat']}",
            f"Tahsil,{h['tahsil']}",
            f"KDV_Oran,%20",
            f"Tevkifat,4/10",
            f"NACE,{NACE}",
        ]
        path.write_text("\n".join(lines) + "\n", encoding="utf-8")


def main():
    wb = Workbook()
    build_kapak(wb)
    build_ozet(wb)
    for no, aciklama, m2 in DONEMLER:
        build_donem_sheet(wb, no, aciklama, m2)

    xlsx = ROOT / "Fatura_Hakedis_Goturu_9400.xlsx"
    wb.save(xlsx)
    export_csv()
    print(f"Yazıldı: {xlsx}")
    for i in range(1, 5):
        print(f"Yazıldı: {ROOT / f'Fatura_Hakedis_Donem_{i}.csv'}")


if __name__ == "__main__":
    main()
