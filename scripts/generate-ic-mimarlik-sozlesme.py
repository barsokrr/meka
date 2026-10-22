#!/usr/bin/env python3
"""Kısa iç mimarlık hizmet sözleşmesi (.docx) üretir."""

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Cm, Pt
from pathlib import Path

OUT = Path(__file__).resolve().parents[1] / "docs" / "Ic-Mimarlik-Hizmet-Sozlesmesi.docx"


def add_heading(doc: Document, text: str, level: int = 1) -> None:
    p = doc.add_heading(text, level=level)
    for run in p.runs:
        run.font.size = Pt(14 if level == 1 else 12)


def add_body(doc: Document, text: str, bold: bool = False) -> None:
    p = doc.add_paragraph()
    run = p.add_run(text)
    run.font.size = Pt(12)
    run.bold = bold
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.line_spacing = 1.25


def add_bullet(doc: Document, text: str) -> None:
    p = doc.add_paragraph(text, style="List Bullet")
    for run in p.runs:
        run.font.size = Pt(12)
    p.paragraph_format.space_after = Pt(4)


def main() -> None:
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Cm(1.5)
    section.bottom_margin = Cm(1.5)
    section.left_margin = Cm(1.8)
    section.right_margin = Cm(1.8)

    style = doc.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(12)

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = title.add_run("İÇ MİMARLIK HİZMET SÖZLEŞMESİ")
    r.bold = True
    r.font.size = Pt(16)
    sub = doc.add_paragraph()
    sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    sr = sub.add_run("Proje çizimi ve şantiye takibi")
    sr.font.size = Pt(12)

    add_body(
        doc,
        "Taslaktır; imzadan önce avukat / İMO mevzuatı kontrolü önerilir.",
        bold=True,
    )

    add_heading(doc, "1. Taraflar", 2)
    add_body(doc, "Müşteri: __________________  Tel: __________")
    add_body(doc, "İç Mimar: __________________  İMO No: ______")
    add_body(doc, "Proje yeri: __________________  Alan: _____ m²")
    add_body(doc, "Tarih: __.__.20__  Sözleşme No: ______")

    add_heading(doc, "2. Konu", 2)
    add_body(
        doc,
        "İç mimar; bu projede çizim ve şantiye takibi (uygulama kontrolü) yapar. "
        "İnşaat işçiliği, statik/elektrik/mekanik proje ve ruhsat işlemleri bu sözleşmeye dahil değildir.",
    )

    add_heading(doc, "3. Hizmetler", 2)
    add_bullet(doc, "Keşif, ölçüm, konsept")
    add_bullet(doc, "Uygulama projeleri (plan, kesit, detay)")
    add_bullet(doc, "Malzeme / renk listesi")
    add_bullet(doc, "Şantiye ziyareti: _____ adet (ziyaret başı _____ saat)")
    add_bullet(doc, "Her ziyaret sonrası kısa tutanak")
    add_body(doc, "Teslim: PDF / DWG — ________________")

    add_heading(doc, "4. Müşteri", 2)
    add_bullet(doc, "Alana erişim ve kararları zamanında vermek")
    add_bullet(doc, "Uygulamayı yürüten firma ile koordinasyon")
    add_bullet(doc, "Bedeli ödeme planına göre ödemek")

    add_heading(doc, "5. Bedel ve ödeme", 2)
    add_body(doc, "Toplam (KDV dahil): __________ TL")
    add_bullet(doc, "Avans %____ — imzada")
    add_bullet(doc, "Konsept onayı %____")
    add_bullet(doc, "Proje teslimi %____")
    add_bullet(doc, "Kapanış %____")
    add_body(doc, "Ödeme: _____ gün içinde havale/EFT. Gecikmede çalışma durabilir.")

    add_heading(doc, "6. Revizyon ve ek iş", 2)
    add_body(
        doc,
        "Her aşamada _____ tur küçük revizyon dahildir. "
        "Plan değişikliği ve ek ziyaret yazılı onay + ayrı ücretlendirilir.",
    )

    add_heading(doc, "7. Süre", 2)
    add_body(doc, "Konsept: __.__.20__  |  Proje: __.__.20__  |  Şantiye: __.__ – __.__.20__")

    add_heading(doc, "8. Telif ve sorumluluk", 2)
    add_body(
        doc,
        "Çizimlerin telif hakkı iç mimara aittir. Ücret tam ödenince müşteri yalnızca bu proje için kullanır. "
        "İç mimar uygulayıcı işçiliğinden sorumlu değildir; mesleki özenle kontrol eder.",
    )

    add_heading(doc, "9. Fesih", 2)
    add_body(
        doc,
        "Taraflar _____ gün önceden yazılı feshedebilir. "
        "Tamamlanan iş ve yapılan ziyaretler oranında bedel tahakkuk eder.",
    )

    add_heading(doc, "10. Uyuşmazlık", 2)
    add_body(doc, "__________ Mahkemeleri yetkilidir. (Tüketici ise özel hükümler saklıdır.)")

    doc.add_paragraph()
    add_body(doc, "Müşteri imza: _________________  Tarih: ______")
    add_body(doc, "İç Mimar imza: _______________  Tarih: ______")

    OUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUT)
    print(OUT)


if __name__ == "__main__":
    main()
